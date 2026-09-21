"""
خروجی‌گیری گزارش در قالب‌های مختلف.

پشتیبانی از پنج قالب: CSV، Excel، JSON، PDF و HTML.

نکته درباره وابستگی‌ها:
    openpyxl و reportlab اختیاری در نظر گرفته شده‌اند. اگر نصب نباشند،
    خطای واضح با راهنمای نصب داده می‌شود و بقیه قالب‌ها همچنان کار
    می‌کنند — نبود یک کتابخانه نباید کل گزارش‌گیری را از کار بیندازد.

نکته درباره فارسی در PDF:
    reportlab به‌طور پیش‌فرض متن راست‌به‌چپ را درست نمی‌چیند. برای اجتناب
    از تولید PDF ناخوانا، عنوان‌ها و برچسب‌های PDF انگلیسی می‌مانند و
    محتوای فارسی در قالب HTML (که مرورگر خودش RTL را مدیریت می‌کند) ارائه
    می‌شود. این یک تصمیم آگاهانه است، نه محدودیت فراموش‌شده.
"""

from __future__ import annotations

import csv
import html
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from app.core.paths import AppPaths, app_paths
from app.exceptions import ValidationError
from app.logging import get_logger
from reports.builder import ReportData

logger = get_logger(__name__)

SUPPORTED_FORMATS = ("csv", "xlsx", "json", "pdf", "html")


class ReportExporter:
    """
    خروجی‌گیر گزارش.

    نمونه‌سازی:
        exporter = ReportExporter()
        path = exporter.export(data, "xlsx")
    """

    def __init__(self, paths: AppPaths | None = None) -> None:
        self._paths = paths or app_paths

    # ------------------------------------------------------------------
    # نقطه ورود
    # ------------------------------------------------------------------
    def export(
        self, data: ReportData, output_format: str, destination: Path | None = None
    ) -> Path:
        """
        خروجی‌گیری گزارش در قالب خواسته‌شده.

        اگر مسیر مقصد داده نشود، فایل در پوشه exports با نام زمان‌دار ساخته
        می‌شود.
        """
        fmt = output_format.lower().strip()
        if fmt in ("excel", "xls"):
            fmt = "xlsx"
        if fmt not in SUPPORTED_FORMATS:
            raise ValidationError(
                f"Unsupported report format: {output_format}",
                details={"supported": list(SUPPORTED_FORMATS)},
            )

        target = destination or self._default_path(data, fmt)
        target.parent.mkdir(parents=True, exist_ok=True)

        writer = {
            "csv": self._write_csv,
            "xlsx": self._write_xlsx,
            "json": self._write_json,
            "pdf": self._write_pdf,
            "html": self._write_html,
        }[fmt]
        writer(data, target)

        logger.info("Report exported: %s (%d rows)", target.name, len(data.rows))
        return target

    def _default_path(self, data: ReportData, fmt: str) -> Path:
        """ساخت نام فایل پیش‌فرض بر پایه عنوان و زمان."""
        self._paths.ensure()
        slug = "".join(c if c.isalnum() else "_" for c in data.title.lower()).strip("_")[:40]
        stamp = data.generated_at.strftime("%Y%m%d_%H%M%S")
        return self._paths.exports_dir / f"{slug}_{stamp}.{fmt}"

    # ------------------------------------------------------------------
    # CSV
    # ------------------------------------------------------------------
    @staticmethod
    def _write_csv(data: ReportData, target: Path) -> None:
        """
        خروجی CSV.

        از `utf-8-sig` استفاده می‌شود تا اکسل در ویندوز متن فارسی را درست
        نمایش دهد؛ بدون BOM، فارسی به‌صورت کاراکترهای درهم دیده می‌شود.
        """
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            if data.is_empty:
                handle.write("No data for the selected period\n")
                return
            fieldnames = list(data.rows[0])
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(data.rows)

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------
    @staticmethod
    def _write_xlsx(data: ReportData, target: Path) -> None:
        """خروجی اکسل با شیت جداگانه برای خلاصه آماری."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:  # pragma: no cover
            raise ValidationError(
                "Excel export requires the 'openpyxl' package. Install it with: pip install openpyxl"
            ) from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Signals"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2F5597")

        if data.rows:
            fieldnames = list(data.rows[0])
            sheet.append(fieldnames)
            for index, _ in enumerate(fieldnames, start=1):
                cell = sheet.cell(row=1, column=index)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            for row in data.rows:
                sheet.append([row.get(name, "") for name in fieldnames])

            # تنظیم عرض ستون‌ها بر پایه بلندترین مقدار
            for index, name in enumerate(fieldnames, start=1):
                longest = max([len(str(name))] + [len(str(r.get(name, ""))) for r in data.rows[:200]])
                sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, 40)
            sheet.freeze_panes = "A2"
        else:
            sheet.append(["No data for the selected period"])

        summary_sheet = workbook.create_sheet("Summary")
        summary_sheet.append(["Report", data.title])
        summary_sheet.append(["Generated at", data.generated_at.strftime("%Y-%m-%d %H:%M:%S UTC")])
        summary_sheet.append(["Period", data.period_label])
        summary_sheet.append([])
        for key, value in data.summary.items():
            summary_sheet.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
        summary_sheet.append([])
        summary_sheet.append(["Disclaimer", "Technical analysis only — not financial advice."])
        summary_sheet.column_dimensions["A"].width = 26
        summary_sheet.column_dimensions["B"].width = 60

        workbook.save(target)

    # ------------------------------------------------------------------
    # JSON
    # ------------------------------------------------------------------
    @staticmethod
    def _write_json(data: ReportData, target: Path) -> None:
        """خروجی JSON کامل، مناسب پردازش ماشینی."""
        target.write_text(
            json.dumps(data.to_dict(), ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------
    @staticmethod
    def _write_pdf(data: ReportData, target: Path) -> None:
        """
        خروجی PDF.

        برچسب‌ها انگلیسی هستند (توضیح در docstring ماژول). برای گزارش
        فارسی خوانا، قالب HTML توصیه می‌شود.
        """
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                Table,
                TableStyle,
            )
        except ImportError as exc:  # pragma: no cover
            raise ValidationError(
                "PDF export requires the 'reportlab' package. Install it with: pip install reportlab"
            ) from exc

        styles = getSampleStyleSheet()
        document = SimpleDocTemplate(
            str(target),
            pagesize=landscape(A4),
            leftMargin=12 * mm, rightMargin=12 * mm,
            topMargin=12 * mm, bottomMargin=12 * mm,
            title=data.title,
        )

        story: list[Any] = [
            Paragraph(data.title, styles["Title"]),
            Paragraph(
                f"{data.app_name} v{data.app_version} — generated "
                f"{data.generated_at.strftime('%Y-%m-%d %H:%M UTC')} — period: {data.period_label}",
                styles["Normal"],
            ),
            Spacer(1, 6 * mm),
        ]

        if data.summary:
            story.append(Paragraph("Summary", styles["Heading2"]))
            summary_rows = [
                [key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)]
                for key, value in data.summary.items()
            ]
            summary_table = Table(summary_rows, colWidths=[50 * mm, 200 * mm])
            summary_table.setStyle(
                TableStyle([
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ])
            )
            story.extend([summary_table, Spacer(1, 6 * mm)])

        story.append(Paragraph("Signals", styles["Heading2"]))
        if data.is_empty:
            story.append(Paragraph("No data for the selected period.", styles["Normal"]))
        else:
            # برای جا شدن در عرض صفحه، فقط ستون‌های کلیدی چاپ می‌شوند
            keys = [
                k for k in ("created_at", "symbol", "direction", "entry_min", "stop_loss",
                            "risk_reward", "leverage", "confidence", "trend", "status")
                if k in data.rows[0]
            ]
            table_data = [[k.replace("_", " ").title() for k in keys]]
            for row in data.rows[:400]:
                table_data.append([str(row.get(k, ""))[:18] for k in keys])

            table = Table(table_data, repeatRows=1)
            table.setStyle(
                TableStyle([
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F5597")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
                ])
            )
            story.append(table)
            if len(data.rows) > 400:
                story.append(Paragraph(f"… and {len(data.rows) - 400} more rows (see CSV/Excel export).", styles["Italic"]))

        story.extend([
            Spacer(1, 8 * mm),
            Paragraph(
                "<b>Disclaimer:</b> This report contains technical analysis only and is not financial "
                "advice. Confidence values express the alignment of analytical factors, not a "
                "probability of profit. No real trades are executed by this application.",
                styles["Italic"],
            ),
        ])
        document.build(story)

    # ------------------------------------------------------------------
    # HTML
    # ------------------------------------------------------------------
    @staticmethod
    def _write_html(data: ReportData, target: Path, *, rtl: bool = False) -> None:
        """
        خروجی HTML مستقل (بدون وابستگی به فایل بیرونی).

        همه سبک‌ها درون فایل قرار می‌گیرند تا گزارش قابل ارسال و باز شدن
        در هر مرورگری باشد.
        """
        direction = "rtl" if rtl else "ltr"
        rows_html = ""
        if data.rows:
            headers = list(data.rows[0])
            head = "".join(f"<th>{html.escape(h.replace('_', ' ').title())}</th>" for h in headers)
            body = "".join(
                "<tr>" + "".join(f"<td>{html.escape(str(row.get(h, '')))}</td>" for h in headers) + "</tr>"
                for row in data.rows
            )
            rows_html = f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"
        else:
            rows_html = "<p class='empty'>No data for the selected period.</p>"

        summary_html = "".join(
            f"<div class='card'><span class='k'>{html.escape(str(key))}</span>"
            f"<span class='v'>{html.escape(json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value))}</span></div>"
            for key, value in data.summary.items()
        )

        target.write_text(
            f"""<!DOCTYPE html>
<html lang="{'fa' if rtl else 'en'}" dir="{direction}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(data.title)}</title>
<style>
  * {{ box-sizing: border-box; }}
  body {{ font-family: "Segoe UI", Tahoma, Arial, sans-serif; margin: 0; padding: 24px;
         background: #f5f6f8; color: #1f2430; }}
  h1 {{ margin: 0 0 4px; font-size: 22px; }}
  .meta {{ color: #61697a; font-size: 13px; margin-bottom: 20px; }}
  .cards {{ display: flex; flex-wrap: wrap; gap: 10px; margin-bottom: 20px; }}
  .card {{ background: #fff; border: 1px solid #e2e5ea; border-radius: 8px; padding: 10px 14px;
           min-width: 150px; }}
  .card .k {{ display: block; font-size: 11px; text-transform: uppercase; color: #7a8296; }}
  .card .v {{ display: block; font-size: 15px; font-weight: 600; margin-top: 3px; }}
  table {{ width: 100%; border-collapse: collapse; background: #fff; font-size: 13px;
           border-radius: 8px; overflow: hidden; box-shadow: 0 1px 3px rgba(0,0,0,.06); }}
  th {{ background: #2f5597; color: #fff; padding: 9px 10px; text-align: {'right' if rtl else 'left'};
        font-weight: 600; white-space: nowrap; }}
  td {{ padding: 8px 10px; border-top: 1px solid #eceef2; }}
  tbody tr:nth-child(even) {{ background: #fafbfc; }}
  .empty {{ background: #fff; padding: 24px; border-radius: 8px; text-align: center; color: #7a8296; }}
  .disclaimer {{ margin-top: 22px; padding: 12px 16px; background: #fff8e6; border: 1px solid #f0dca6;
                 border-radius: 8px; font-size: 12px; color: #6b5a20; line-height: 1.7; }}
</style>
</head>
<body>
  <h1>{html.escape(data.title)}</h1>
  <div class="meta">{html.escape(data.app_name)} v{html.escape(data.app_version)} &middot;
    generated {data.generated_at.strftime('%Y-%m-%d %H:%M UTC')} &middot; period: {html.escape(data.period_label)}</div>
  <div class="cards">{summary_html}</div>
  {rows_html}
  <div class="disclaimer">
    <b>Disclaimer:</b> This report contains technical analysis only and is <b>not financial advice</b>.
    Confidence values express how well the analytical factors agree with each other &mdash; they are
    <b>not</b> a probability of profit. This application does not execute real trades.
  </div>
</body>
</html>""",
            encoding="utf-8",
        )

    def export_all(self, data: ReportData) -> dict[str, Path]:
        """خروجی‌گیری هم‌زمان در همه قالب‌ها (برای آزمون و پشتیبان‌گیری گزارش)."""
        results: dict[str, Path] = {}
        for fmt in SUPPORTED_FORMATS:
            try:
                results[fmt] = self.export(data, fmt)
            except ValidationError as exc:
                logger.warning("Skipping %s export: %s", fmt, exc.message)
        return results
