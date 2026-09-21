"""آزمون گزارش‌گیری در همه قالب‌ها."""

from __future__ import annotations

import json
import zipfile
from datetime import UTC, datetime

import pytest

from app.core.constants import AnalysisStatus, MarketStructureType, SignalDirection, TrendDirection
from app.core.models import TradingSignal
from app.core.paths import AppPaths
from app.database.repositories import SignalRepository
from app.database.session import DatabaseManager
from app.exceptions import ValidationError
from reports import ReportBuilder, ReportExporter


@pytest.fixture()
def report_setup(temp_paths: AppPaths):
    """پایگاه داده با چند سیگنال نمونه."""
    database = DatabaseManager(temp_paths.database_url)
    database.create_all()
    repository = SignalRepository(database)

    for index in range(6):
        direction = [SignalDirection.LONG, SignalDirection.SHORT, SignalDirection.WAIT][index % 3]
        is_wait = direction == SignalDirection.WAIT
        repository.save_signal(
            TradingSignal(
                symbol="BTC/USDT" if index % 2 else "ETH/USDT",
                exchange="test",
                direction=direction,
                entry_min=None if is_wait else 100.0,
                entry_max=None if is_wait else 101.0,
                stop_loss=None if is_wait else 98.0,
                take_profits=[] if is_wait else [104.0, 108.0],
                risk_reward=None if is_wait else 2.0,
                leverage=3,
                confidence=50 + index,
                trend=TrendDirection.BULLISH,
                market_structure=MarketStructureType.RANGING,
                reason="آزمون گزارش با متن فارسی",
                invalidation="close below stop",
                timeframes=["4h", "1h"],
                indicators_used=["RSI"],
                status=AnalysisStatus.OK,
                created_at=datetime.now(UTC),
            )
        )
    return temp_paths, ReportBuilder(repository)


def test_report_contains_all_signals(report_setup) -> None:
    """گزارش باید همه سیگنال‌های بازه را در بر بگیرد."""
    _, builder = report_setup
    data = builder.build_signal_report(days=30)
    assert len(data.rows) == 6
    assert data.summary["total"] == 6


def test_summary_has_no_win_rate(report_setup) -> None:
    """
    نسخه ۱ معامله واقعی ندارد، پس ادعای نرخ برد گمراه‌کننده است و نباید
    در خلاصه وجود داشته باشد.
    """
    _, builder = report_setup
    summary = builder.build_signal_report(days=30).summary
    assert not any("win" in key.lower() or "success" in key.lower() for key in summary)


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json", "pdf", "html"])
def test_every_format_produces_a_file(report_setup, fmt: str) -> None:
    """هر پنج قالب باید فایل غیرخالی بسازند."""
    paths, builder = report_setup
    data = builder.build_signal_report(days=30)
    path = ReportExporter(paths).export(data, fmt)
    assert path.exists() and path.stat().st_size > 0


def test_json_export_is_parsable(report_setup) -> None:
    """خروجی JSON باید قابل تجزیه و شامل سلب مسئولیت باشد."""
    paths, builder = report_setup
    data = builder.build_signal_report(days=30)
    path = ReportExporter(paths).export(data, "json")
    payload = json.loads(path.read_text(encoding="utf-8"))
    assert len(payload["rows"]) == 6
    assert "not financial advice" in payload["disclaimer"]


def test_excel_has_two_sheets(report_setup) -> None:
    """فایل اکسل باید شیت داده و شیت خلاصه داشته باشد."""
    from openpyxl import load_workbook

    paths, builder = report_setup
    data = builder.build_signal_report(days=30)
    path = ReportExporter(paths).export(data, "xlsx")
    assert zipfile.is_zipfile(path)
    assert load_workbook(path).sheetnames == ["Signals", "Summary"]


def test_pdf_has_valid_header(report_setup) -> None:
    """فایل PDF باید امضای درست داشته باشد."""
    paths, builder = report_setup
    data = builder.build_signal_report(days=30)
    path = ReportExporter(paths).export(data, "pdf")
    assert path.read_bytes().startswith(b"%PDF")


def test_html_is_self_contained(report_setup) -> None:
    """
    گزارش HTML باید مستقل باشد؛ وابستگی به فایل بیرونی یعنی گزارش
    ارسال‌شده برای دیگران خراب دیده می‌شود.
    """
    paths, builder = report_setup
    data = builder.build_signal_report(days=30)
    content = ReportExporter(paths).export(data, "html").read_text(encoding="utf-8")
    assert "<style>" in content
    assert "http://" not in content and "https://" not in content


def test_csv_uses_bom_for_excel_compatibility(report_setup) -> None:
    """
    بدون BOM، اکسل ویندوز متن فارسی را درهم نشان می‌دهد.
    """
    paths, builder = report_setup
    data = builder.build_signal_report(days=30)
    path = ReportExporter(paths).export(data, "csv")
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")


def test_empty_report_does_not_crash(report_setup) -> None:
    """گزارش خالی باید بدون خطا ساخته شود."""
    paths, builder = report_setup
    data = builder.build_signal_report(days=30, symbol="NOTHING/USDT")
    assert data.is_empty
    for fmt in ("csv", "xlsx", "json", "pdf", "html"):
        assert ReportExporter(paths).export(data, fmt).exists()


def test_unsupported_format_is_rejected(report_setup) -> None:
    """قالب پشتیبانی‌نشده باید خطای واضح بدهد."""
    paths, builder = report_setup
    data = builder.build_signal_report(days=30)
    with pytest.raises(ValidationError):
        ReportExporter(paths).export(data, "docx")
